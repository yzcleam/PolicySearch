from __future__ import annotations

import hashlib
import html
import io
import json
import re
import time
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any, Callable, Iterable, Iterator

import ijson
from openpyxl import load_workbook

from ..models import Document
from . import catalog
from .index_service import index_service, to_source

PKU_LAW_COLUMNS = [
    "id",
    "law_db_name",
    "title",
    "detail_url",
    "province",
    "city",
    "county",
    "category_1",
    "category_2",
    "law_type",
    "pub_depart",
    "pub_num",
    "pub_date",
    "use_date",
    "is_time",
    "detail_flag",
    "detail_html",
    "detail_md",
    "add_time",
]

TITLE_KEYS = ("title", "标题")
DOC_ID_KEYS = ("doc_id", "document_id", "id", "ID", "source_id")
DOC_NUMBER_KEYS = ("doc_number", "pub_num", "文号", "document_number")
ISSUER_KEYS = ("issuer", "pub_depart", "发布机关", "发文机关")
REGION_KEYS = ("region", "province", "地区", "省份")
CATEGORY_KEYS = ("category", "category_1", "政策类型", "主题分类")
TAG_KEYS = ("tags", "标签", "主题标签")
ISSUE_DATE_KEYS = ("issue_date", "pub_date", "发布日期", "date")
EFFECTIVE_DATE_KEYS = ("effective_date", "use_date", "实施日期", "生效日期")
CONTENT_KEYS = ("content", "detail_md", "detail_flag", "detail_html", "正文", "text")
SUMMARY_KEYS = ("summary", "摘要", "abstract")
SOURCE_URL_KEYS = ("source_url", "detail_url", "url", "链接")


@dataclass
class ImportResult:
    source: str = ""
    format: str = "auto"
    total: int = 0
    imported: int = 0
    indexed: int = 0
    skipped: int = 0
    failed: int = 0
    errors: list[str] = field(default_factory=list)
    message: str = ""
    elapsed_seconds: float = 0.0


ProgressCallback = Callable[[ImportResult], None]


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bytes):
        return value.decode("utf-8", errors="replace").strip()
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value).strip()


def _clean_content(value: Any) -> str:
    text = _text(value)
    if not text:
        return ""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = text.replace("\u00a0", " ")
    return text.strip()


def _strip_html(value: Any) -> str:
    text = _clean_content(value)
    if "<" not in text or ">" not in text:
        return text
    text = re.sub(r"(?is)<(script|style).*?</\1>", " ", text)
    text = re.sub(r"(?i)<br\s*/?>", "\n", text)
    text = re.sub(r"(?i)</p\s*>", "\n\n", text)
    text = re.sub(r"(?s)<[^>]+>", " ", text)
    text = html.unescape(text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _first(record: dict[str, Any], keys: Iterable[str]) -> Any:
    for key in keys:
        if key in record and record[key] not in (None, ""):
            return record[key]
    return None


def _split_values(value: Any) -> list[str]:
    if value is None:
        return []
    if isinstance(value, list):
        values = value
    else:
        values = re.split(r"[;,；、|]", str(value))
    return [str(item).strip() for item in values if str(item).strip()]


def parse_date(value: Any) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    text = text.replace("年", "-").replace("月", "-").replace("日", "")
    text = text.replace("/", "-").replace(".", "-")
    text = re.sub(r"\s+.*$", "", text)
    for fmt in ("%Y-%m-%d", "%Y-%m", "%Y%m%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _status(value: Any, time_status: Any) -> str:
    raw = _text(value)
    if raw:
        return raw[:32]
    marker = _text(time_status)
    if "现行有效" in marker or "有效" in marker:
        return "active"
    if any(item in marker for item in ("失效", "废止", "已被修改")):
        return "inactive"
    return (marker or "active")[:32]


def normalize_record(raw: Any, source_format: str = "json") -> dict[str, Any] | None:
    if not isinstance(raw, dict):
        return None
    record = {str(key).strip().lower(): value for key, value in raw.items()}

    title = _clean_content(_first(record, TITLE_KEYS))
    if not title:
        return None
    title = title[:512]

    raw_id = _first(record, DOC_ID_KEYS)
    source_url = _text(_first(record, SOURCE_URL_KEYS))
    if raw_id not in (None, ""):
        doc_id = str(raw_id).strip()
        if doc_id.isdigit():
            prefix = "pkulaw" if "law_db_name" in record or "detail_url" in record else "import"
            doc_id = f"{prefix}-{doc_id}"
    else:
        digest = hashlib.sha1(f"{source_url}|{title}".encode("utf-8")).hexdigest()[:24]
        doc_id = f"import-{digest}"
    if len(doc_id) > 64:
        doc_id = doc_id[:64]

    content = ""
    for key in CONTENT_KEYS:
        candidate = record.get(key)
        if candidate:
            content = _strip_html(candidate)
            if content:
                break
    if not content:
        content = title

    summary_value = _first(record, SUMMARY_KEYS)
    summary = _strip_html(summary_value) if summary_value else ""
    if not summary:
        summary = re.sub(r"\s+", " ", content)[:300]

    category_value = _first(record, CATEGORY_KEYS)
    category_parts = _split_values(category_value)
    category = (category_parts[0] if category_parts else "")[:64]

    tags: list[str] = []
    raw_tags = _first(record, TAG_KEYS)
    if isinstance(raw_tags, list):
        tags.extend(str(item).strip() for item in raw_tags if str(item).strip())
    elif raw_tags:
        tags.extend(_split_values(raw_tags))
    for key in ("category_1", "category_2", "law_type", "law_db_name", "city", "county"):
        tags.extend(_split_values(record.get(key)))
    tags = list(dict.fromkeys(item for item in tags if item))[:50]

    issue_value = _first(record, ISSUE_DATE_KEYS)
    effective_value = _first(record, EFFECTIVE_DATE_KEYS)

    region_value = _first(record, REGION_KEYS)
    if not region_value:
        region_value = record.get("city") or record.get("county")
    region = _text(region_value)[:64]

    issuer = _text(_first(record, ISSUER_KEYS))[:128]
    doc_number = _text(_first(record, DOC_NUMBER_KEYS))[:128]
    file_hash = _text(record.get("file_hash"))
    if not file_hash:
        file_hash = hashlib.sha256(
            f"{doc_id}|{title}|{content}".encode("utf-8")
        ).hexdigest()

    issue_date = parse_date(issue_value)
    effective_date = parse_date(effective_value)
    law_db_name = _text(record.get("law_db_name"))[:64]
    category_1 = _text(record.get("category_1"))
    category_2 = _text(record.get("category_2"))
    law_type = _text(record.get("law_type"))[:128]
    province = _text(record.get("province"))[:512] or region
    city = _text(record.get("city"))[:64]
    county = _text(record.get("county"))[:64]
    is_time = _text(record.get("is_time"))[:64]
    detail_flag = _clean_content(record.get("detail_flag"))
    detail_html = _clean_content(record.get("detail_html"))
    detail_md = _clean_content(record.get("detail_md")) or content
    add_time = _text(record.get("add_time"))[:64]

    return {
        "doc_id": doc_id,
        "title": title,
        "doc_number": doc_number,
        "issuer": issuer,
        "region": region,
        "category": category,
        "tags": tags,
        "issue_date": issue_date,
        "effective_date": effective_date,
        "content": content,
        "summary": summary,
        "source_url": source_url[:512],
        "file_hash": file_hash[:64],
        "status": _status(record.get("status"), record.get("is_time")),
        "law_db_name": law_db_name,
        "detail_url": source_url[:512],
        "province": province,
        "city": city,
        "county": county,
        "category_1": category_1,
        "category_2": category_2,
        "law_type": law_type,
        "pub_depart": issuer,
        "pub_num": doc_number,
        "pub_date": issue_date,
        "use_date": effective_date,
        "is_time": is_time,
        "detail_flag": detail_flag,
        "detail_html": detail_html,
        "detail_md": detail_md,
        "add_time": add_time,
    }


# --------------------------------------------------------------------------- JSON


def _iter_json_stream(stream: Any) -> Iterator[dict[str, Any]]:
    prefix_bytes = stream.read(65536)
    stream.seek(0)
    prefix = prefix_bytes.decode("utf-8-sig", errors="replace").lstrip()
    if prefix.startswith("["):
        path = "item"
    elif '"RECORDS"' in prefix or '"records"' in prefix:
        path = "RECORDS.item" if '"RECORDS"' in prefix else "records.item"
    elif '"data"' in prefix:
        path = "data.item"
    else:
        raise ValueError(
            "JSON 顶层必须是数组，或包含 RECORDS / records / data 数组的对象"
        )
    for item in ijson.items(stream, path, use_float=True):
        if isinstance(item, dict):
            yield item


def _iter_json_lines(stream: Any) -> Iterator[dict[str, Any]]:
    for line in io.TextIOWrapper(stream, encoding="utf-8-sig", errors="replace"):
        line = line.strip()
        if not line:
            continue
        item = json.loads(line)
        if isinstance(item, dict):
            yield item


# ---------------------------------------------------------------------------- SQL


def _iter_sql_statements(stream: Any) -> Iterator[str]:
    text_stream = io.TextIOWrapper(stream, encoding="utf-8-sig", errors="replace")
    buffer: list[str] = []
    in_single = in_double = in_backtick = False
    escaped = False
    for chunk in iter(lambda: text_stream.read(1024 * 1024), ""):
        for char in chunk:
            buffer.append(char)
            if escaped:
                escaped = False
                continue
            if in_single and char == "\\":
                escaped = True
                continue
            if char == "'" and not in_double and not in_backtick:
                in_single = not in_single
            elif char == '"' and not in_single and not in_backtick:
                in_double = not in_double
            elif char == "`" and not in_single and not in_double:
                in_backtick = not in_backtick
            elif char == ";" and not in_single and not in_double and not in_backtick:
                statement = "".join(buffer).strip()
                buffer.clear()
                if statement:
                    yield statement
            if len(buffer) > 128 * 1024 * 1024:
                raise ValueError("单条 SQL 语句超过 128MB，无法安全解析")
    if buffer:
        statement = "".join(buffer).strip()
        if statement:
            yield statement


def _split_top_level(value: str, separator: str = ",") -> Iterator[str]:
    start = 0
    depth = 0
    in_single = in_double = in_backtick = False
    escaped = False
    for index, char in enumerate(value):
        if escaped:
            escaped = False
            continue
        if in_single and char == "\\":
            escaped = True
            continue
        if char == "'" and not in_double and not in_backtick:
            in_single = not in_single
        elif char == '"' and not in_single and not in_backtick:
            in_double = not in_double
        elif char == "`" and not in_single and not in_double:
            in_backtick = not in_backtick
        elif not in_single and not in_double and not in_backtick:
            if char == "(":
                depth += 1
            elif char == ")":
                depth -= 1
            elif char == separator and depth == 0:
                yield value[start:index]
                start = index + 1
    yield value[start:]


def _unescape_sql(value: str) -> str:
    body = value[1:-1]
    body = body.replace("''", "'").replace('""', '"')
    replacements = {
        "0": "\0",
        "b": "\b",
        "n": "\n",
        "r": "\r",
        "t": "\t",
        "Z": "\x1a",
        "\\": "\\",
        "'": "'",
        '"': '"',
        "%": "%",
        "_": "_",
    }
    return re.sub(r"\\(.)", lambda match: replacements.get(match.group(1), match.group(1)), body)


def _parse_sql_value(value: str) -> Any:
    token = value.strip()
    if not token:
        return ""
    if token.upper() == "NULL":
        return None
    if token.upper() == "TRUE":
        return True
    if token.upper() == "FALSE":
        return False
    if token[0] in ("'", '"') and token[-1] == token[0]:
        return _unescape_sql(token)
    if re.fullmatch(r"-?\d+", token):
        try:
            return int(token)
        except ValueError:
            pass
    if re.fullmatch(r"-?\d+\.\d+", token):
        try:
            return float(token)
        except ValueError:
            pass
    return token


def _iter_sql_tuples(values: str) -> Iterator[str]:
    index = 0
    length = len(values)
    while index < length:
        while index < length and values[index] in " \t\r\n,":
            index += 1
        if index >= length or values[index] != "(":
            break
        start = index + 1
        depth = 1
        in_single = in_double = in_backtick = False
        escaped = False
        index += 1
        while index < length and depth:
            char = values[index]
            if escaped:
                escaped = False
            elif in_single and char == "\\":
                escaped = True
            elif char == "'" and not in_double and not in_backtick:
                in_single = not in_single
            elif char == '"' and not in_single and not in_backtick:
                in_double = not in_double
            elif char == "`" and not in_single and not in_double:
                in_backtick = not in_backtick
            elif not in_single and not in_double and not in_backtick:
                if char == "(":
                    depth += 1
                elif char == ")":
                    depth -= 1
            index += 1
        if depth:
            raise ValueError("SQL VALUES 中存在未闭合的元组")
        yield values[start : index - 1]


def _iter_sql_stream(stream: Any) -> Iterator[dict[str, Any]]:
    insert_re = re.compile(
        r"(?is)^\s*INSERT\s+(?:IGNORE\s+)?INTO\s+([^\s(]+)\s*(\([^)]*\))?\s*VALUES\s*(.*)$"
    )
    for statement in _iter_sql_statements(stream):
        match = insert_re.match(statement)
        if not match:
            continue
        table = match.group(1).replace("`", "")
        columns_text = match.group(2)
        values_text = match.group(3).strip()
        if columns_text:
            columns = [
                part.strip().strip("`").lower()
                for part in _split_top_level(columns_text.strip()[1:-1])
                if part.strip()
            ]
        elif "pkulaw_content" in table.lower():
            columns = PKU_LAW_COLUMNS
        else:
            columns = PKU_LAW_COLUMNS
        for tuple_text in _iter_sql_tuples(values_text):
            values = [
                _parse_sql_value(part)
                for part in _split_top_level(tuple_text)
            ]
            if len(values) < len(columns):
                values.extend([None] * (len(columns) - len(values)))
            yield dict(zip(columns, values))


# --------------------------------------------------------------------------- Excel


def _iter_excel(source: Any, sheet: str | None = None) -> Iterator[dict[str, Any]]:
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        worksheet = workbook[sheet] if sheet else workbook.worksheets[0]
        rows = worksheet.iter_rows(values_only=True)
        headers: list[str] | None = None
        for row in rows:
            if row is None or not any(value is not None and str(value).strip() for value in row):
                continue
            if headers is None:
                headers = [
                    str(value).strip().lower() if value is not None else ""
                    for value in row
                ]
                continue
            if not headers:
                continue
            yield {
                header: value
                for header, value in zip(headers, row)
                if header and value is not None
            }
    finally:
        workbook.close()


# --------------------------------------------------------------------------- ZIP


def _detect_format(path: Path, requested: str) -> str:
    if requested != "auto":
        return requested
    suffix = path.suffix.lower()
    if suffix in {".json", ".jsonl"}:
        return "json"
    if suffix in {".xlsx", ".xlsm"}:
        return "excel"
    if suffix == ".sql":
        return "sql"
    if suffix == ".zip":
        with zipfile.ZipFile(path) as archive:
            names = archive.namelist()
        for member in names:
            lower = member.lower()
            if lower.endswith((".json", ".jsonl")):
                return "json"
            if lower.endswith(".sql"):
                return "sql"
            if lower.endswith((".xlsx", ".xlsm")):
                return "excel"
        raise ValueError("ZIP 中没有找到 .json、.sql 或 .xlsx 数据文件")
    raise ValueError(f"无法识别导入格式: {path.suffix or path.name}")


def _zip_member(archive: zipfile.ZipFile, fmt: str) -> str:
    endings = {
        "json": (".json", ".jsonl"),
        "sql": (".sql",),
        "excel": (".xlsx", ".xlsm"),
    }[fmt]
    members = [name for name in archive.namelist() if not name.endswith("/") and name.lower().endswith(endings)]
    if not members:
        raise ValueError(f"ZIP 中没有找到 {fmt} 数据文件")
    if len(members) > 1:
        raise ValueError(f"ZIP 中有多个 {fmt} 数据文件，请解压后指定具体文件")
    return members[0]


def _open_zip_member(path: Path, fmt: str, password: str | None) -> Any:
    archive = zipfile.ZipFile(path)
    member = _zip_member(archive, fmt)
    try:
        return archive, member, archive.open(member, pwd=password.encode("utf-8") if password else None)
    except RuntimeError as exc:
        archive.close()
        message = str(exc).lower()
        if "encrypted" in message or "password" in message:
            raise ValueError(
                "ZIP 已加密，请提供 zip_password / --zip-password"
            ) from exc
        raise
    except Exception:
        archive.close()
        raise


# ------------------------------------------------------------------------ public


def iter_records(
    path: str | Path,
    fmt: str = "auto",
    zip_password: str | None = None,
    sheet: str | None = None,
) -> Iterator[dict[str, Any]]:
    source = Path(path)
    if not source.exists():
        raise FileNotFoundError(f"文件不存在: {source}")
    resolved = _detect_format(source, fmt)

    if source.suffix.lower() == ".zip" and resolved != "excel":
        archive, member, stream = _open_zip_member(source, resolved, zip_password)
        try:
            if resolved == "json" and member.lower().endswith(".jsonl"):
                yield from _iter_json_lines(stream)
            elif resolved == "json":
                yield from _iter_json_stream(stream)
            elif resolved == "sql":
                yield from _iter_sql_stream(stream)
        finally:
            stream.close()
            archive.close()
        return

    if source.suffix.lower() == ".zip" and resolved == "excel":
        archive = zipfile.ZipFile(source)
        member = _zip_member(archive, "excel")
        try:
            data = archive.read(
                member,
                pwd=zip_password.encode("utf-8") if zip_password else None,
            )
        finally:
            archive.close()
        yield from _iter_excel(io.BytesIO(data), sheet)
        return

    if resolved == "json":
        with source.open("rb") as stream:
            if source.suffix.lower() == ".jsonl":
                yield from _iter_json_lines(stream)
            else:
                yield from _iter_json_stream(stream)
    elif resolved == "sql":
        with source.open("rb") as stream:
            yield from _iter_sql_stream(stream)
    elif resolved == "excel":
        yield from _iter_excel(source, sheet)
    else:
        raise ValueError(f"不支持的格式: {resolved}")


def import_source(
    path: str | Path,
    fmt: str = "auto",
    batch_size: int = 500,
    limit: int = 0,
    index: bool = True,
    reset: bool = False,
    zip_password: str | None = None,
    sheet: str | None = None,
    progress: ProgressCallback | None = None,
) -> ImportResult:
    started = time.perf_counter()
    result = ImportResult(source=str(path), format=fmt)
    if batch_size < 1 or batch_size > 5000:
        raise ValueError("batch_size 必须在 1-5000 之间")
    resolved = _detect_format(Path(path), fmt)
    result.format = resolved

    if reset:
        catalog.clear_all()
        if index:
            index_service.delete_all()
    if index:
        index_service.ensure_index()

    batch: list[dict[str, Any]] = []

    def flush() -> None:
        nonlocal batch
        if not batch:
            return
        result.imported += catalog.upsert_documents(batch)
        if index:
            documents = [to_source(Document(**item)) for item in batch]
            result.indexed += index_service.index_documents(documents, refresh=False)
        batch = []
        if progress:
            progress(result)

    for raw in iter_records(path, resolved, zip_password=zip_password, sheet=sheet):
        result.total += 1
        failed_record = False
        try:
            normalized = normalize_record(raw, resolved)
        except Exception as exc:
            result.failed += 1
            failed_record = True
            if len(result.errors) < 20:
                result.errors.append(f"记录 {result.total}: {exc}")
            normalized = None
        if normalized is None:
            if not failed_record:
                result.skipped += 1
        else:
            batch.append(normalized)
            if len(batch) >= batch_size:
                flush()
        if limit and result.total >= limit:
            break

    flush()
    if index:
        try:
            index_service.client.indices.refresh(index=index_service.index)
        except Exception:
            pass
    result.elapsed_seconds = round(time.perf_counter() - started, 2)
    result.message = (
        f"{result.format}: 读取 {result.total} 条，导入 {result.imported} 条，"
        f"索引 {result.indexed} 条，跳过 {result.skipped} 条，失败 {result.failed} 条"
    )
    if progress:
        progress(result)
    return result
